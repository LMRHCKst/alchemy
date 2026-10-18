# -*- coding: utf-8 -*-
"""Тесты выгрузки в DOCX."""

from __future__ import annotations

import os
import tempfile
import unittest

from alchemy.core import report as R
from alchemy.engines import balancer as BAL
from alchemy.engines import thermo as TH


HAS_DOCX = R.available()


@unittest.skipUnless(HAS_DOCX, "python-docx не установлен")
class TestResultExport(unittest.TestCase):
    """Выгрузка одного результата."""

    def setUp(self):
        self.directory = tempfile.mkdtemp(prefix="alchemy_test_")
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        for name in os.listdir(self.directory):
            try:
                os.remove(os.path.join(self.directory, name))
            except OSError:
                pass
        try:
            os.rmdir(self.directory)
        except OSError:
            pass

    def test_file_created(self):
        result = TH.reaction_enthalpy("CaCO3(тв) = CaO(тв) + CO2")
        path = os.path.join(self.directory, "result.docx")
        R.export_result(path, result)
        self.assertTrue(os.path.exists(path))
        self.assertGreater(os.path.getsize(path), 1000)

    def test_content_has_value(self):
        from docx import Document

        result = TH.reaction_enthalpy("CaCO3(тв) = CaO(тв) + CO2")
        path = os.path.join(self.directory, "result.docx")
        R.export_result(path, result)

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        cells = [cell.text for table in document.tables
                 for row in table.rows for cell in row.cells]
        combined = text + "\n" + "\n".join(cells)

        self.assertIn("178", combined)
        self.assertIn("кДж/моль", combined)

    def test_latex_included_when_source_given(self):
        from docx import Document

        result = TH.reaction_enthalpy("CaCO3(тв) = CaO(тв) + CO2")
        path = os.path.join(self.directory, "result.docx")
        R.export_result(path, result, source="CaCO3(тв) = CaO(тв) + CO2")

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("LaTeX", text)
        self.assertIn(r"\rightarrow", text)

    def test_steps_present(self):
        from docx import Document

        result = TH.reaction_enthalpy("CaCO3(тв) = CaO(тв) + CO2")
        path = os.path.join(self.directory, "result.docx")
        R.export_result(path, result)

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Ход расчёта", text)

    def test_warnings_present(self):
        from docx import Document

        result = TH.reaction_enthalpy("CaCO3(тв) = CaO(тв) + CO2")
        result.with_warning("Проверочное предупреждение")
        path = os.path.join(self.directory, "result.docx")
        R.export_result(path, result)

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Проверочное предупреждение", text)


@unittest.skipUnless(HAS_DOCX, "python-docx не установлен")
class TestJournalExport(unittest.TestCase):
    """Выгрузка журнала."""

    def setUp(self):
        self.directory = tempfile.mkdtemp(prefix="alchemy_test_")
        self.addCleanup(self._cleanup)

    def _cleanup(self):
        for name in os.listdir(self.directory):
            try:
                os.remove(os.path.join(self.directory, name))
            except OSError:
                pass
        try:
            os.rmdir(self.directory)
        except OSError:
            pass

    def _entries(self):
        thermo = TH.reaction_enthalpy("CaCO3(тв) = CaO(тв) + CO2")
        balanced = BAL.balance("2H2 + O2 = H2O")
        return [
            {"time": "10:00:01", "action": "Термохимия",
             "result": thermo.formatted(), "detail": thermo.summary(),
             "warnings": []},
            {"time": "10:00:02", "action": "Уравнивание",
             "result": balanced.formatted(), "detail": balanced.summary(),
             "warnings": []},
        ]

    def test_empty_journal(self):
        from docx import Document

        path = os.path.join(self.directory, "empty.docx")
        R.export_journal(path, [])
        self.assertTrue(os.path.exists(path))
        text = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("Журнал пуст", text)

    def test_entries_table(self):
        from docx import Document

        path = os.path.join(self.directory, "journal.docx")
        R.export_journal(path, self._entries())

        document = Document(path)
        self.assertGreaterEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual(len(table.rows), 3)      # шапка + две записи

    def test_titles_present(self):
        from docx import Document

        entries = self._entries()
        path = os.path.join(self.directory, "journal.docx")
        R.export_journal(path, entries)

        text = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("Термохимия", text)
        self.assertIn("Уравнивание", text)


class TestHelpers(unittest.TestCase):
    """Вспомогательные функции."""

    def test_suggested_name(self):
        name = R.suggested_name()
        self.assertTrue(name.startswith("alchemy_"))
        self.assertTrue(name.endswith(".docx"))

    def test_latex_of_prefers_source(self):
        from alchemy.core.results import Result

        result = Result(value=1.0, unit="г", label="Масса")
        self.assertIn(r"\mathrm", R.latex_of(result, source="CuSO4"))

    def test_latex_of_empty_without_source(self):
        from alchemy.core.results import Result

        result = Result(value=1.0, unit="г", label="Масса")
        self.assertEqual(R.latex_of(result), "")

    def test_available(self):
        self.assertIsInstance(R.available(), bool)


if __name__ == "__main__":
    unittest.main()
